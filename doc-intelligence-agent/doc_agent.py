#!/usr/bin/env python3
"""
Document Intelligence Agent — skenuje složku OneDrive, analyzuje dokumenty přes Claude API
a ukládá výsledky do Supabase tabulky doc_intelligence.
"""

from dotenv import load_dotenv

load_dotenv()

import hashlib
import json
import os
import sys
import traceback
from datetime import datetime, timedelta, timezone
from pathlib import Path

import anthropic
from supabase import Client, create_client

SUPABASE_URL = os.getenv("SUPABASE_URL", "").strip()
SUPABASE_KEY = os.getenv("SUPABASE_KEY", "").strip()
ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY", "").strip()
CLICKUP_API_KEY = os.getenv("CLICKUP_API_KEY", "").strip()
CLICKUP_LIST_ID = os.getenv("CLICKUP_LIST_ID", "901514038952").strip()
WATCHED_FOLDER = os.getenv("WATCHED_FOLDER", "").strip()
HOURS_BACK = int(os.getenv("HOURS_BACK", "24"))

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".txt", ".md", ".rtf",
    ".xlsx", ".xls", ".pptx", ".ppt", ".odt"
}

CATEGORIES = [
    "Granty a projekty", "Administrativa", "Výzkum", "Výuka", "Personalistika",
    "Smlouvy", "Zprávy a analýzy", "Komunikace", "Ostatní"
]

LOG_DIR = Path(__file__).resolve().parent / "logs"
LOG_FILE = LOG_DIR / "agent.log"


def log(msg: str) -> None:
    line = msg.rstrip()
    print(line)
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(line + "\n")


def require_env() -> None:
    missing = []
    if not SUPABASE_URL:
        missing.append("SUPABASE_URL")
    if not SUPABASE_KEY:
        missing.append("SUPABASE_KEY")
    if not ANTHROPIC_API_KEY:
        missing.append("ANTHROPIC_API_KEY")
    if not WATCHED_FOLDER:
        missing.append("WATCHED_FOLDER")
    if missing:
        raise RuntimeError("Chybí proměnné v .env: " + ", ".join(missing))
    if not Path(WATCHED_FOLDER).is_dir():
        raise RuntimeError(f"Složka neexistuje: {WATCHED_FOLDER}")


def get_supabase() -> Client:
    return create_client(SUPABASE_URL, SUPABASE_KEY)


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _extract_doc_with_word(file_path: str) -> str:
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc = None
    try:
        doc = word.Documents.Open(file_path)
        return doc.Content.Text or ""
    finally:
        if doc is not None:
            doc.Close(False)
        word.Quit()


def _extract_doc_with_antiword(file_path: str) -> str:
    import subprocess

    result = subprocess.run(
        ["antiword", file_path],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=60,
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError((result.stderr or result.stdout or "antiword selhal").strip())
    return result.stdout


def extract_text_preview(path: Path, max_chars: int = 12000) -> str:
    file_path = str(path.resolve())
    ext = path.suffix.lower()
    try:
        if ext in {".txt", ".md", ".rtf"}:
            return path.read_text(encoding="utf-8", errors="replace")[:max_chars]
        if ext == ".docx":
            from docx import Document

            doc = Document(file_path)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())[:max_chars]
        if ext == ".doc":
            try:
                text = _extract_doc_with_word(file_path)
            except Exception:
                try:
                    text = _extract_doc_with_antiword(file_path)
                except Exception as exc:
                    return f"[Nepodařilo se extrahovat text z .doc: {exc}]"
            return text[:max_chars]
        if ext == ".xlsx":
            import openpyxl

            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            try:
                text = ""
                for sheet in wb.worksheets[:3]:
                    text += f"[List: {sheet.title}]\n"
                    for row in sheet.iter_rows(max_row=50, values_only=True):
                        line = " | ".join(str(v) for v in row if v is not None)
                        if line.strip():
                            text += line + "\n"
                    if len(text) > max_chars:
                        break
                return text[:max_chars]
            finally:
                wb.close()
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(file_path)
            parts = []
            for page in reader.pages[:20]:
                parts.append(page.extract_text() or "")
            return "\n".join(parts)[:max_chars]
    except Exception as exc:
        return f"[Nepodařilo se extrahovat text: {exc}]"
    return f"[Nepodporovaný formát pro extrakci: {ext}]"


def extract_text(path: Path) -> str:
    return extract_text_preview(path, max_chars=12000)


def find_recent_files(root: Path, hours_back: int) -> list[Path]:
    cutoff = datetime.now(timezone.utc) - timedelta(hours=hours_back)
    results = []
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        if path.name.startswith("~$") or path.name.startswith("."):
            continue
        mtime = datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc)
        if mtime >= cutoff:
            results.append(path)
    return sorted(results, key=lambda p: p.stat().st_mtime, reverse=True)


def hash_exists(sb: Client, digest: str) -> bool:
    res = (
        sb.table("doc_intelligence")
        .select("id")
        .eq("file_hash", digest)
        .limit(1)
        .execute()
    )
    return bool(res.data)


def analyze_document(ai: anthropic.Anthropic, text: str, filename: str) -> dict:
    prompt = f"""Analyzuj tento dokument pro asistentku prorektorky UHK pro vědu a výzkum.

Soubor: {filename}

Obsah (výňatek):
{text[:8000]}

Odpověz POUZE validním JSON objektem (bez markdown) s klíči:
- tema (stručný název/téma, max 120 znaků)
- souhrn (2-4 věty česky)
- kategorie (jedna z: {", ".join(CATEGORIES)})
- dulezitost (celé číslo 1-5, kde 5 = kritické)
- klicova_slova (pole max 8 řetězců)
- akce_doporucena (jedna konkrétní doporučená akce nebo prázdný řetězec)"""

    response = ai.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=900,
        messages=[{"role": "user", "content": prompt}],
    )
    raw = response.content[0].text.strip()
    if raw.startswith("```"):
        raw = raw.split("```", 2)[1]
        if raw.startswith("json"):
            raw = raw[4:]
        raw = raw.strip()
    data = json.loads(raw)
    kat = data.get("kategorie", "Ostatní")
    if kat not in CATEGORIES:
        kat = "Ostatní"
    dulezitost = int(data.get("dulezitost", 3))
    dulezitost = max(1, min(5, dulezitost))
    keywords = data.get("klicova_slova") or []
    if not isinstance(keywords, list):
        keywords = [str(keywords)]
    return {
        "tema": str(data.get("tema", filename))[:200],
        "souhrn": str(data.get("souhrn", ""))[:2000],
        "kategorie": kat,
        "dulezitost": dulezitost,
        "klicova_slova": [str(k)[:60] for k in keywords[:8]],
        "akce_doporucena": str(data.get("akce_doporucena", ""))[:500],
    }


def save_document(sb: Client, path: Path, root: Path, analysis: dict, digest: str) -> dict:
    rel = str(path.relative_to(root))
    stat = path.stat()
    row = {
        "file_name": path.name,
        "file_path": str(path),
        "file_url": path.as_uri(),
        "relative_path": rel,
        "folder": str(path.parent.relative_to(root)) if path.parent != root else "",
        "extension": path.suffix.lower().lstrip("."),
        "size_kb": round(stat.st_size / 1024, 1),
        "file_modified_at": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
        "file_hash": digest,
        "stav": "nový",
        **analysis,
    }
    res = sb.table("doc_intelligence").insert(row).execute()
    return res.data[0] if res.data else row


def generate_daily_summary(ai: anthropic.Anthropic, saved_docs: list) -> str:
    if not saved_docs:
        return "Dnes nebyly přidány žádné nové dokumenty."

    doc_list = "\n".join(
        f"- {d.get('tema', d.get('file_name', '?'))} "
        f"({d.get('kategorie', '?')}, priorita {d.get('dulezitost', '?')}/5): "
        f"{(d.get('souhrn') or '')[:100]}…"
        for d in saved_docs
    )

    response = ai.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=600,
        messages=[{
            "role": "user",
            "content": f"""Jsi asistentka prorektorky UHK pro vědu a výzkum.
Napiš stručný denní přehled (max 5 vět) nových dokumentů:

{doc_list}

Upozorni na kritické dokumenty. Piš česky, osobně, bez nadpisů."""
        }],
    )
    return response.content[0].text


def run_agent() -> int:
    require_env()
    root = Path(WATCHED_FOLDER)
    sb = get_supabase()
    ai = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)

    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    log("=" * 60)
    log(f"Document Intelligence Agent — {now}")
    log(f"Složka: {root}")
    log("=" * 60)

    files = find_recent_files(root, HOURS_BACK)
    log(f"Nalezeno {len(files)} nových/změněných souborů za posledních {HOURS_BACK}h")

    saved = []
    for i, path in enumerate(files, 1):
        rel = str(path.relative_to(root))
        log(f"[{i}/{len(files)}] {rel}")
        try:
            digest = file_hash(path)
            if hash_exists(sb, digest):
                log("  → Přeskočeno (již v databázi)")
                continue
            log("  → Analyzuji přes Claude API...")
            text = extract_text(path)
            analysis = analyze_document(ai, text, path.name)
            log(
                f"  → Téma: {analysis['tema']} | "
                f"Kategorie: {analysis['kategorie']} | "
                f"Důležitost: {analysis['dulezitost']}/5"
            )
            record = save_document(sb, path, root, analysis, digest)
            saved.append(record)
            log(f"  ✓ Uloženo (id: {record.get('id', '?')})")
        except Exception as exc:
            log(f"  ✗ Chyba: {exc}")
            traceback.print_exc()

    if saved:
        summary = generate_daily_summary(ai, saved)
        log("")
        log("Denní souhrn:")
        log(summary)
        try:
            sb.table("doc_intelligence_summary").insert({
                "summary_text": summary,
                "doc_count": len(saved),
                "created_at": datetime.now(timezone.utc).isoformat(),
            }).execute()
            log("  ✓ Denní souhrn uložen do doc_intelligence_summary")
        except Exception as exc:
            log(f"  ✗ Nepodařilo se uložit souhrn: {exc}")

    log("")
    log(f"Hotovo — uloženo {len(saved)} nových dokumentů.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(run_agent())
    except Exception as err:
        log(f"FATÁLNÍ CHYBA: {err}")
        traceback.print_exc()
        sys.exit(1)
